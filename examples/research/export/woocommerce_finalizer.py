"""
WooCommerce CSV Finalizer

Converts normalized wine data (JSON) into two WooCommerce-importable CSVs
compliant with bilingual ES/EN requirements.
"""

import json
import csv
import re
from pathlib import Path
from typing import Dict, List, Any, Tuple, Union
from decimal import Decimal

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils.dataframe import dataframe_to_rows

    EXCEL_AVAILABLE = True
except ImportError:
    EXCEL_AVAILABLE = False


class WooCommerceFinalizer:
    """Main class for converting wine data to WooCommerce CSV format."""

    def __init__(self, mapping_dir: str):
        self.mapping_dir = Path(mapping_dir)
        self._country_mappings = {}
        self._type_mappings = {}
        self._primary_grapes_mappings = {}
        self._shortlist_grapes = set()
        self._load_mappings()

    def _load_mappings(self):
        """Load canonical attribute mappings from CSV files."""
        # Load countries
        countries_file = self.mapping_dir / "countries.csv"
        if countries_file.exists():
            with open(countries_file, "r", encoding="utf-8") as f:
                reader = csv.DictReader(f)
                for row in reader:
                    es_val = row["ES"].strip()
                    en_val = row["EN"].strip()
                    self._country_mappings[es_val] = (es_val, en_val)

        # Load types
        types_file = self.mapping_dir / "types.csv"
        if types_file.exists():
            with open(types_file, "r", encoding="utf-8") as f:
                reader = csv.DictReader(f)
                for row in reader:
                    es_val = row["ES"].strip()
                    en_val = row["EN"].strip()
                    self._type_mappings[es_val] = (es_val, en_val)

        # Load primary grapes
        grapes_file = self.mapping_dir / "primary_grapes.csv"
        if grapes_file.exists():
            with open(grapes_file, "r", encoding="utf-8") as f:
                reader = csv.DictReader(f)
                for row in reader:
                    es_val = row["ES"].strip()
                    en_val = row["EN"].strip()
                    is_shortlist = row.get("shortlist", "false").lower() == "true"
                    self._primary_grapes_mappings[es_val] = (es_val, en_val)
                    if is_shortlist:
                        self._shortlist_grapes.add(es_val)

    def map_country(self, raw: str) -> Tuple[str, str]:
        """Map country to ES/EN canonical values."""
        raw = raw.strip()
        if raw in self._country_mappings:
            return self._country_mappings[raw]
        # Default fallback - return as-is for both languages
        return (raw, raw)

    def map_type(self, raw: str) -> Tuple[str, str]:
        """Map wine type to ES/EN canonical values."""
        raw = raw.strip()
        if raw in self._type_mappings:
            return self._type_mappings[raw]
        # Default fallback
        return (raw, raw)

    def map_primary_grapes(self, raw_list: List[str]) -> Tuple[List[str], List[str]]:
        """Map primary grapes to ES/EN with shortlist constraint."""
        es_grapes = []
        en_grapes = []

        for grape in raw_list:
            grape = grape.strip()
            if grape in self._primary_grapes_mappings:
                es_val, en_val = self._primary_grapes_mappings[grape]
                if grape in self._shortlist_grapes:
                    es_grapes.append(es_val)
                    en_grapes.append(en_val)

        # If no shortlist grapes found, add "Other varietals"
        if not es_grapes:
            es_grapes.append("Otros varietales")
            en_grapes.append("Other varietals")

        return (es_grapes, en_grapes)

    def normalize_abv(
        self, value: Union[str, float, int], style: str = "percent"
    ) -> str:
        """Normalize alcohol by volume to specified style."""
        if value is None:
            return ""

        # Convert to string and replace comma with dot
        str_val = str(value).replace(",", ".")

        try:
            # Parse as decimal for precision
            decimal_val = Decimal(str_val)
            if style == "percent":
                return f"{decimal_val}%"
            return str(decimal_val)
        except:
            return str(value)

    def normalize_temp(self, temp_str: str, style: str = "14–17 °C") -> str:
        """Normalize serving temperature to specified style."""
        if not temp_str:
            return ""

        # Extract numbers from temperature string
        numbers = re.findall(r"\d+", temp_str)
        if len(numbers) >= 2:
            if style == "14–17 °C":
                return f"{numbers[0]}–{numbers[1]} °C"
            elif style == "14º–17ºC":
                return f"{numbers[0]}º–{numbers[1]}ºC"
        elif len(numbers) == 1:
            if style == "14–17 °C":
                return f"{numbers[0]} °C"
            elif style == "14º–17ºC":
                return f"{numbers[0]}ºC"

        return temp_str

    def strip_vintage(self, name: str) -> str:
        """Remove vintage years from wine name."""
        if not name:
            return ""

        # Remove 4-digit years (vintages)
        result = re.sub(r"\b(19|20)\d{2}\b", "", name)
        # Clean up extra spaces
        result = re.sub(r"\s+", " ", result).strip()
        return result

    def format_price(self, price: Union[str, float, int]) -> str:
        """Format price to dot decimal string."""
        if price is None:
            return "0.00"

        try:
            # Convert to decimal for precision
            decimal_val = Decimal(str(price).replace(",", "."))
            return f"{decimal_val:.2f}"
        except:
            return "0.00"

    def build_es_row(
        self, wine: Dict[str, Any], include_attribute_data: bool = False
    ) -> Dict[str, str]:
        """Build ES CSV row from wine data."""
        row = {}

        # Basic fields
        row["sku"] = wine.get("sku", "")
        row["post_title"] = self.strip_vintage(
            wine.get("name_es", wine.get("normalized_name", ""))
        )
        row["post_content"] = wine.get("desc_es", "")
        row["post_excerpt"] = wine.get("tasting_es", "")
        row["regular_price"] = self.format_price(wine.get("price"))

        # Stock status
        stock = wine.get("stock")
        stock_status = wine.get("stock_status")
        if stock_status:
            row["stock_status"] = (
                "instock" if stock_status == "instock" else "outofstock"
            )
        elif isinstance(stock, bool):
            row["stock_status"] = "instock" if stock else "outofstock"
        else:
            row["stock_status"] = "instock"

        # Attributes
        country_es, _ = self.map_country(wine.get("country", ""))
        row["attribute:pa_pais"] = country_es
        row["attribute:pa_region"] = wine.get("region_es", wine.get("region", ""))

        type_es, _ = self.map_type(wine.get("type", ""))
        row["attribute:pa_tipo"] = type_es

        # Build grape blend text
        grapes = wine.get("grapes", [])
        if wine.get("blend_text"):
            row["attribute:pa_uva"] = wine["blend_text"]
        elif grapes:
            grape_parts = []
            for grape in grapes:
                name = grape.get("name", "")
                percent = grape.get("percent")
                if percent:
                    grape_parts.append(f"{name} {percent}%")
                else:
                    grape_parts.append(name)
            row["attribute:pa_uva"] = "|".join(grape_parts)
        else:
            row["attribute:pa_uva"] = ""

        # Primary grapes (shortlist only)
        primary_grapes = wine.get(
            "primary_grapes", [grape.get("name", "") for grape in grapes]
        )
        es_primary, _ = self.map_primary_grapes(primary_grapes)
        row["attribute:pa_uva-principal"] = "|".join(es_primary)

        row["attribute:pa_enologo-bodega"] = wine.get(
            "winery_es", wine.get("producer", "")
        )
        row["attribute:pa_como-servirlo"] = self.normalize_temp(
            wine.get("serving_temp_es", "")
        )
        row["attribute:pa_graduacion-alcoholica"] = self.normalize_abv(
            wine.get("alcohol_by_volume")
        )

        # Optional closure type
        if wine.get("closure_type"):
            row["attribute:pa_tipo-tapon"] = wine["closure_type"]

        # Attribute positioning data if requested
        if include_attribute_data:
            positions = {
                "pa_maridaje": 1,
                "pa_region": 2,
                "pa_enologo-bodega": 3,
                "pa_graduacion-alcoholica": 4,
                "pa_tipo-tapon": 5,
                "pa_pais": 6,
                "pa_tipo": 7,
                "pa_uva-principal": 8,
                "pa_como-servirlo": 9,
                "pa_uva": 10,
            }
            for attr, pos in positions.items():
                if f"attribute:{attr}" in row and row[f"attribute:{attr}"]:
                    row[f"attribute_data:{attr}"] = f"{pos}|1|0"

        return row

    def build_en_row(
        self,
        wine: Dict[str, Any],
        es_sku: str,
        en_sku: str,
        include_attribute_data: bool = False,
    ) -> Dict[str, str]:
        """Build EN CSV row from wine data."""
        # Start with ES row as base
        row = self.build_es_row(wine, include_attribute_data)

        # Update with EN-specific values
        row["sku"] = en_sku
        row["wpml:original_product_sku"] = es_sku

        # Use EN text fields if available, otherwise keep ES
        if wine.get("name_en"):
            row["post_title"] = self.strip_vintage(wine["name_en"])
        if wine.get("desc_en"):
            row["post_content"] = wine["desc_en"]
        if wine.get("tasting_en"):
            row["post_excerpt"] = wine["tasting_en"]

        # Translate attributes to EN
        _, country_en = self.map_country(wine.get("country", ""))
        row["attribute:pa_pais"] = country_en

        if wine.get("region_en"):
            row["attribute:pa_region"] = wine["region_en"]

        _, type_en = self.map_type(wine.get("type", ""))
        row["attribute:pa_tipo"] = type_en

        # Primary grapes in EN
        grapes = wine.get("grapes", [])
        primary_grapes = wine.get(
            "primary_grapes", [grape.get("name", "") for grape in grapes]
        )
        _, en_primary = self.map_primary_grapes(primary_grapes)
        row["attribute:pa_uva-principal"] = "|".join(en_primary)

        if wine.get("winery_en"):
            row["attribute:pa_enologo-bodega"] = wine["winery_en"]
        if wine.get("serving_temp_en"):
            row["attribute:pa_como-servirlo"] = self.normalize_temp(
                wine["serving_temp_en"]
            )

        return row

    def validate_wine(self, wine: Dict[str, Any]) -> List[str]:
        """Validate a single wine record and return list of errors."""
        errors = []

        required_fields = [
            "normalized_name",
            "producer",
            "region",
            "appellation",
            "vintage",
            "grapes",
        ]
        for field in required_fields:
            if not wine.get(field):
                errors.append(f"Missing required field: {field}")

        # Validate grapes structure
        grapes = wine.get("grapes", [])
        if grapes and not isinstance(grapes, list):
            errors.append("'grapes' must be a list")
        elif grapes:
            for i, grape in enumerate(grapes):
                if not isinstance(grape, dict) or not grape.get("name"):
                    errors.append(f"Grape at index {i} must have 'name' field")

        return errors

    def write_excel_file(
        self,
        rows: List[Dict[str, str]],
        columns: List[str],
        output_path: str,
        sheet_name: str = "Products",
    ) -> bool:
        """Write data to Excel file with formatting."""
        if not EXCEL_AVAILABLE:
            return False

        try:
            wb = Workbook()
            ws = wb.active
            ws.title = sheet_name

            # Write header with formatting
            header_font = Font(bold=True, color="FFFFFF")
            header_fill = PatternFill(
                start_color="366092", end_color="366092", fill_type="solid"
            )
            header_alignment = Alignment(horizontal="center", vertical="center")

            for col_idx, column in enumerate(columns, 1):
                cell = ws.cell(row=1, column=col_idx, value=column)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = header_alignment

            # Write data rows
            for row_idx, row_data in enumerate(rows, 2):
                for col_idx, column in enumerate(columns, 1):
                    value = row_data.get(column, "")

                    # Convert complex data structures to strings
                    if isinstance(value, (dict, list)):
                        if isinstance(value, dict):
                            # Convert dict to readable string
                            value = "; ".join(
                                [f"{k}: {v}" for k, v in value.items() if v]
                            )
                        else:
                            # Convert list to string
                            value = "; ".join([str(item) for item in value if item])

                    # Ensure value is string and not too long for Excel
                    value = str(value)[:32767] if value else ""  # Excel cell limit

                    cell = ws.cell(row=row_idx, column=col_idx, value=value)

                    # Format price columns
                    if column == "regular_price" and value:
                        try:
                            cell.value = float(value)
                            cell.number_format = "0.00"
                        except ValueError:
                            pass

            # Auto-adjust column widths
            for column in ws.columns:
                max_length = 0
                column_letter = column[0].column_letter

                for cell in column:
                    try:
                        if len(str(cell.value)) > max_length:
                            max_length = len(str(cell.value))
                    except:
                        pass

                adjusted_width = min(max_length + 2, 50)  # Cap at 50 characters
                ws.column_dimensions[column_letter].width = adjusted_width

            # Freeze header row
            ws.freeze_panes = "A2"

            wb.save(output_path)
            return True

        except Exception as e:
            print(f"❌ Failed to write Excel file: {e}")
            return False


def finalize_wine_to_woocommerce_csv(
    data_path: str,
    mapping_dir: str,
    output_es: str = "products_es.csv",
    output_en: str = "products_en.csv",
    sku_strategy: str = "suffix",  # one of {"suffix","same"}
    include_attribute_data: bool = False,
    normalize_temp_style: str = "14–17 °C",
    normalize_abv_style: str = "percent",
    output_format: str = "csv",  # one of {"csv", "excel", "both"}
) -> dict:
    """
    Convert normalized wine data to WooCommerce CSV format.

    Returns a report dict with keys:
      - wines_processed (int)
      - rows_written_es (int)
      - rows_written_en (int)
      - warnings (list[str])
      - errors (list[str])  # empty if successful
    """

    report = {
        "wines_processed": 0,
        "rows_written_es": 0,
        "rows_written_en": 0,
        "warnings": [],
        "errors": [],
    }

    try:
        # Load input data
        with open(data_path, "r", encoding="utf-8") as f:
            data = json.load(f)

        # Handle both single object and array
        if isinstance(data, dict):
            wines = [data]
        elif isinstance(data, list):
            wines = data
        else:
            report["errors"].append("Input data must be a JSON object or array")
            return report

        # Initialize finalizer
        finalizer = WooCommerceFinalizer(mapping_dir)

        # Validate all wines first
        valid_wines = []
        for i, wine in enumerate(wines):
            validation_errors = finalizer.validate_wine(wine)
            if validation_errors:
                for error in validation_errors:
                    report["errors"].append(f"Wine {i}: {error}")
            else:
                valid_wines.append(wine)

        # If any validation errors, stop processing
        if report["errors"]:
            return report

        # Define column order
        base_columns = [
            "sku",
            "post_title",
            "post_content",
            "post_excerpt",
            "regular_price",
            "stock_status",
            "attribute:pa_pais",
            "attribute:pa_region",
            "attribute:pa_tipo",
            "attribute:pa_uva",
            "attribute:pa_uva-principal",
            "attribute:pa_enologo-bodega",
            "attribute:pa_como-servirlo",
            "attribute:pa_graduacion-alcoholica",
        ]

        optional_columns = ["attribute:pa_tipo-tapon"]

        # Add attribute data columns if requested
        if include_attribute_data:
            attr_data_columns = [
                "attribute_data:pa_maridaje",
                "attribute_data:pa_region",
                "attribute_data:pa_enologo-bodega",
                "attribute_data:pa_graduacion-alcoholica",
                "attribute_data:pa_tipo-tapon",
                "attribute_data:pa_pais",
                "attribute_data:pa_tipo",
                "attribute_data:pa_uva-principal",
                "attribute_data:pa_como-servirlo",
                "attribute_data:pa_uva",
            ]
            base_columns.extend(attr_data_columns)

        en_columns = ["wpml:original_product_sku"] + base_columns

        # Process wines and build rows
        es_rows = []
        en_rows = []

        for wine in valid_wines:
            # Generate SKU if not provided
            if not wine.get("sku"):
                # Simple SKU generation - could be made more sophisticated
                wine["sku"] = f"wine_{len(es_rows) + 1:04d}"

            es_sku = wine["sku"]
            en_sku = es_sku + "-en" if sku_strategy == "suffix" else es_sku

            # Build rows
            es_row = finalizer.build_es_row(wine, include_attribute_data)
            en_row = finalizer.build_en_row(
                wine, es_sku, en_sku, include_attribute_data
            )

            es_rows.append(es_row)
            en_rows.append(en_row)

        # Determine output files
        es_csv_path = output_es
        en_csv_path = output_en
        es_excel_path = output_es.replace(".csv", ".xlsx")
        en_excel_path = output_en.replace(".csv", ".xlsx")

        # Write CSV files if requested
        if output_format in ["csv", "both"]:
            # Write ES CSV
            with open(es_csv_path, "w", encoding="utf-8", newline="") as f:
                writer = csv.DictWriter(
                    f,
                    fieldnames=base_columns + optional_columns,
                    delimiter=";",
                    extrasaction="ignore",
                )
                writer.writeheader()
                for row in es_rows:
                    writer.writerow(row)

            # Write EN CSV
            with open(en_csv_path, "w", encoding="utf-8", newline="") as f:
                writer = csv.DictWriter(
                    f,
                    fieldnames=en_columns + optional_columns,
                    delimiter=";",
                    extrasaction="ignore",
                )
                writer.writeheader()
                for row in en_rows:
                    writer.writerow(row)

        # Write Excel files if requested
        if output_format in ["excel", "both"]:
            if not EXCEL_AVAILABLE:
                report["warnings"].append(
                    "Excel output requested but openpyxl not available"
                )
            else:
                # Write ES Excel
                if finalizer.write_excel_file(
                    es_rows,
                    base_columns + optional_columns,
                    es_excel_path,
                    "Products ES",
                ):
                    report["excel_files_written"] = (
                        report.get("excel_files_written", 0) + 1
                    )
                else:
                    report["warnings"].append("Failed to write ES Excel file")

                # Write EN Excel
                if finalizer.write_excel_file(
                    en_rows, en_columns + optional_columns, en_excel_path, "Products EN"
                ):
                    report["excel_files_written"] = (
                        report.get("excel_files_written", 0) + 1
                    )
                else:
                    report["warnings"].append("Failed to write EN Excel file")

        # Update report
        report["wines_processed"] = len(valid_wines)
        report["rows_written_es"] = len(es_rows)
        report["rows_written_en"] = len(en_rows)

    except Exception as e:
        report["errors"].append(f"Unexpected error: {str(e)}")

    return report


if __name__ == "__main__":
    # Example usage
    result = finalize_wine_to_woocommerce_csv(
        data_path="wine_data.json",
        mapping_dir="mappings/",
        output_es="products_es.csv",
        output_en="products_en.csv",
    )
    print(f"Processed {result['wines_processed']} wines")
    if result["errors"]:
        print("Errors:", result["errors"])
    if result["warnings"]:
        print("Warnings:", result["warnings"])
